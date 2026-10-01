"""One-off loader: DGEG weekly Excel -> BigQuery (national history).

Run as `python -m src.load_history` (add `--dry-run` to parse without loading).
Details of the file layout are in docs/sources.md; table layout is in docs/schema.md.

Two raw tables are rebuilt on every run (WRITE_TRUNCATE): the Excel is a full-history
snapshot, so reloading is idempotent and never duplicates weeks.
- raw_weekly_national: pivot cache (2010 onwards), week x fuel x component.
- raw_history_2004_2009: flat sheet, week x fuel (retail price only).
"""

import argparse
import io
import logging
import os
import re
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime
from typing import Any, Iterable, Iterator

import requests
from dotenv import load_dotenv
from google.cloud import bigquery

from src.collect import build_client

DEFAULT_EXCEL_URL = "https://www.dgeg.gov.pt/media/x13gkhm4/dgeg-pcr-2004-2026_40_pt.xlsx"
USER_AGENT = "fuel-monitor-portfolio"  # no personal data here
REQUEST_TIMEOUT_S = 120

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
CACHE_DEFINITION = "xl/pivotCache/pivotCacheDefinition1.xml"
CACHE_RECORDS = "xl/pivotCache/pivotCacheRecords1.xml"
FLAT_SHEET_MARKER = "2004 a 2009"  # the real sheet name has a cedilla; match on the period
FLAT_HEADER_LABEL = "data"
ANNUAL_AVERAGE_PREFIX = "media anual"
FLAT_DATE_PATTERN = re.compile(r"^\s*(\d{4}-\d{2}-\d{2})")  # tolerates footnote marks like "(*)"

WEEKLY_TABLE = "raw_weekly_national"
LEGACY_TABLE = "raw_history_2004_2009"

WEEKLY_SCHEMA = [
    bigquery.SchemaField("week_start", "DATE"),
    bigquery.SchemaField("fuel_type", "STRING"),
    bigquery.SchemaField("component", "STRING"),
    bigquery.SchemaField("value_eur", "FLOAT"),
]
LEGACY_SCHEMA = [
    bigquery.SchemaField("week_start", "DATE"),
    bigquery.SchemaField("fuel_type", "STRING"),
    bigquery.SchemaField("value_eur", "FLOAT"),
]

log = logging.getLogger("load_history")


def normalise(text: str) -> str:
    """Lower-case, accent-free, trimmed: lets us match DGEG labels without caring about encoding."""
    decomposed = unicodedata.normalize("NFKD", text)
    return "".join(c for c in decomposed if not unicodedata.combining(c)).strip().lower()


def local_name(element: ET.Element) -> str:
    """Tag without its XML namespace."""
    return element.tag.rsplit("}", 1)[-1]


def download_excel(url: str) -> bytes:
    """TODO 1: fetch the workbook; raise if the response is not a zip (e.g. DGEG moved the file)."""
    response = requests.get(url, headers={"User-Agent": USER_AGENT}, timeout=REQUEST_TIMEOUT_S)
    response.raise_for_status()
    if not zipfile.is_zipfile(io.BytesIO(response.content)):
        raise RuntimeError(f"{url} did not return an .xlsx file: the DGEG media id probably changed")
    log.info("downloaded %d bytes from %s", len(response.content), url)
    return response.content


def _item_value(element: ET.Element) -> Any:
    """Value of one pivot item (`<s>`, `<n>`, `<d>`, `<b>`); `<m/>` is a missing value."""
    tag = local_name(element)
    if tag == "m":
        return None
    raw = element.attrib["v"]
    if tag == "n":
        return float(raw)
    if tag == "d":
        return datetime.fromisoformat(raw).date()
    return raw


def read_shared_items(definition_xml: bytes) -> dict[str, list[Any]]:
    """Field name (normalised) -> its shared-items list, in on-disk order (records are positional)."""
    root = ET.fromstring(definition_xml)
    fields: dict[str, list[Any]] = {}
    for field in root.find("m:cacheFields", NS):
        name = normalise(field.attrib["name"])
        if name in fields:
            raise RuntimeError(f"duplicate pivot field name: {name!r}")
        shared = field.find("m:sharedItems", NS)
        fields[name] = [_item_value(item) for item in shared] if shared is not None else []
    return fields


def iter_pivot_records(records_xml: bytes, shared: dict[str, list[Any]]) -> Iterator[dict[str, Any]]:
    """TODO 2: yield one dict per pivot record, with `<x v=idx/>` references resolved to values."""
    field_names = list(shared)
    for _, element in ET.iterparse(io.BytesIO(records_xml), events=("end",)):
        if local_name(element) != "r":
            continue
        children = list(element)
        if len(children) != len(field_names):
            raise RuntimeError(f"record has {len(children)} cells, expected {len(field_names)}")
        record: dict[str, Any] = {}
        for name, child in zip(field_names, children):
            if local_name(child) == "x":
                record[name] = shared[name][int(child.attrib["v"])]
            else:
                record[name] = _item_value(child)
        yield record
        element.clear()


def weekly_national_rows(workbook: zipfile.ZipFile) -> list[dict[str, Any]]:
    """TODO 3: long table week_start, fuel_type, component, value_eur from the pivot cache."""
    definition_xml = workbook.read(CACHE_DEFINITION)
    expected = int(ET.fromstring(definition_xml).attrib["recordCount"])
    shared = read_shared_items(definition_xml)
    records = list(iter_pivot_records(workbook.read(CACHE_RECORDS), shared))
    if len(records) != expected:
        raise RuntimeError(f"pivot cache has {len(records)} records but recordCount={expected}")
    if any(r["data"] is None or r["valor medio"] is None for r in records):
        raise RuntimeError("pivot cache contains records with a missing date or value")
    return [
        {
            "week_start": r["data"].isoformat(),
            "fuel_type": r["combustivel"],
            "component": r["componente"],
            "value_eur": r["valor medio"],
        }
        for r in records
    ]


def parse_flat_date(cell: Any) -> date | None:
    """A real date, or text starting with an ISO date (DGEG adds footnote marks); None otherwise."""
    if isinstance(cell, datetime):
        return cell.date()
    if isinstance(cell, str):
        match = FLAT_DATE_PATTERN.match(cell)
        if match:
            return date.fromisoformat(match.group(1))
    return None


def legacy_rows(sheet_rows: Iterable[tuple]) -> list[dict[str, Any]]:
    """TODO 4: flat sheet -> long table (retail price per fuel); annual-average rows are dropped."""
    iterator = iter(sheet_rows)
    header: list[str | None] = []
    for row in iterator:
        if row and isinstance(row[0], str) and normalise(row[0]) == FLAT_HEADER_LABEL:
            header = [str(c).strip() if c is not None else None for c in row]
            break
    if not header:
        raise RuntimeError("header row ('Data') not found in the 2004-2009 sheet")

    rows: list[dict[str, Any]] = []
    for row in iterator:
        first = row[0] if row else None
        if first is None:
            continue
        if isinstance(first, str) and normalise(first).startswith(ANNUAL_AVERAGE_PREFIX):
            continue
        week = parse_flat_date(first)
        if week is None:
            raise RuntimeError(f"unrecognised row label in the 2004-2009 sheet: {first!r}")
        for fuel, value in zip(header[1:], row[1:]):
            if fuel is None or value is None:
                continue
            rows.append({"week_start": week.isoformat(), "fuel_type": fuel, "value_eur": float(value)})
    return rows


def legacy_sheet_rows(excel_bytes: bytes) -> Iterator[tuple]:
    """Values of the 2004-2009 sheet (openpyxl is only needed here; the pivot cache is read raw)."""
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(excel_bytes), read_only=True, data_only=True)
    matches = [name for name in workbook.sheetnames if FLAT_SHEET_MARKER in name]
    if len(matches) != 1:
        raise RuntimeError(f"expected one sheet containing {FLAT_SHEET_MARKER!r}, found {matches}")
    yield from workbook[matches[0]].iter_rows(values_only=True)


def replace_table(client: bigquery.Client, table: str, schema: list[bigquery.SchemaField], rows: list[dict]) -> None:
    """TODO 5: load job that replaces the table content; waits and raises on failure."""
    table_id = f"{client.project}.{os.environ['BQ_DATASET']}.{table}"
    job_config = bigquery.LoadJobConfig(
        schema=schema,
        write_disposition=bigquery.WriteDisposition.WRITE_TRUNCATE,
        create_disposition=bigquery.CreateDisposition.CREATE_IF_NEEDED,
        ignore_unknown_values=False,
    )
    client.load_table_from_json(rows, table_id, job_config=job_config).result()
    log.info("loaded %d rows into %s", len(rows), table_id)


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    load_dotenv()
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--dry-run", action="store_true", help="parse and report, do not touch BigQuery")
    args = parser.parse_args()

    excel_bytes = download_excel(os.environ.get("DGEG_EXCEL_URL", DEFAULT_EXCEL_URL))
    with zipfile.ZipFile(io.BytesIO(excel_bytes)) as workbook:
        weekly = weekly_national_rows(workbook)
    legacy = legacy_rows(legacy_sheet_rows(excel_bytes))
    log.info("parsed %d weekly rows and %d legacy rows", len(weekly), len(legacy))
    if not weekly or not legacy:
        raise RuntimeError("a parsed table is empty: refusing to load")

    if args.dry_run:
        return
    client = build_client()
    replace_table(client, WEEKLY_TABLE, WEEKLY_SCHEMA, weekly)
    replace_table(client, LEGACY_TABLE, LEGACY_SCHEMA, legacy)


if __name__ == "__main__":
    main()
